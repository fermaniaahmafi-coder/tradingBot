import os
import json
import time
import subprocess
from flask import Flask, jsonify, request, render_template, send_file
from config import Config
from db_manager import get_wallet, set_wallet, get_positions, save_positions, get_trades, get_recent_activities, reset_db
from bot_logger import log_event, get_recent_logs
from scanner import get_token_details

app = Flask(__name__)

BOT_STATE_FILE = "bot_state.json"

def get_bot_state():
    default_state = {
        "mode": "testing",      # "testing" or "production"
        "bot_status": "running" # "running", "paused", "stopped"
    }
    if os.path.exists(BOT_STATE_FILE):
        try:
            with open(BOT_STATE_FILE, "r") as f:
                return {**default_state, **json.load(f)}
        except Exception:
            pass
    return default_state

def save_bot_state(state):
    try:
        with open(BOT_STATE_FILE, "w") as f:
            json.dump(state, f, indent=2)
    except Exception:
        pass

def load_data_with_fallback(mode=None):
    """Load positions, history, cash, and logs according to active mode."""
    state = get_bot_state()
    if mode is None:
        mode = state.get("mode", "testing")

    if mode == "production":
        portfolio_file = "portfolio_prod.json"
        history_file = "trade_history_prod.json"
        activities_file = "activities_prod.json"
        default_cash = 250.00
    else:
        portfolio_file = getattr(Config, "DATA_FILE", "portfolio.json")
        history_file = getattr(Config, "HISTORY_FILE", "trade_history.json")
        activities_file = "activities.json"
        default_cash = 100.00

    # In testing mode on Linux VPS, check if SQLite DB has active data
    if mode == "testing":
        try:
            db_path = getattr(Config, "DB_PATH", "/home/trading/trading.db")
            if os.path.exists(db_path):
                db_cash = get_wallet()
                db_positions = get_positions()
                db_history = get_trades(500)
                db_logs = get_recent_activities(40)
                if db_history or db_positions or db_cash > 0:
                    for idx, t in enumerate(db_history, 1):
                        if not t.get("id"):
                            t["id"] = idx
                    return db_cash, db_positions, db_history, db_logs
        except Exception:
            pass

    cash = default_cash
    positions = []
    history = []
    logs = []

    if os.path.exists(portfolio_file):
        try:
            with open(portfolio_file, "r") as f:
                port_data = json.load(f)
                cash = float(port_data.get("cash", default_cash))
                positions = port_data.get("positions", [])
        except Exception:
            pass

    if os.path.exists(history_file):
        try:
            with open(history_file, "r") as f:
                history = json.load(f)
                for idx, t in enumerate(history, 1):
                    if not t.get("id"):
                        t["id"] = idx
        except Exception:
            pass

    if os.path.exists(activities_file):
        try:
            with open(activities_file, "r") as f:
                logs = json.load(f)[:40]
        except Exception:
            pass
    elif mode == "testing" and os.path.exists("activities.json"):
        try:
            with open("activities.json", "r") as f:
                logs = json.load(f)[:40]
        except Exception:
            pass

    return cash, positions, history, logs


@app.route("/")
def index():
    return render_template("dashboard.html")


@app.route("/analytics")
def analytics_page():
    return render_template("analytics.html")



@app.route("/transactions")
def transactions_page():
    return render_template("transactions.html")


@app.route("/coin/<address>")
def coin_detail(address):
    return render_template("coin.html")


@app.route("/api/status")
def api_status():
    state = get_bot_state()
    mode = request.args.get("mode", state.get("mode", "testing")).lower()
    bot_status = state.get("bot_status", "running")
    # Check real systemd service status on Linux VPS
    try:
        check = subprocess.run(["systemctl", "is-active", "trading-bot"], capture_output=True, text=True, timeout=1)
        if check.returncode == 0:
            if check.stdout.strip() == "active":
                bot_status = "running"
            elif bot_status != "paused":
                bot_status = "stopped"
    except Exception:
        pass

    cash, positions, history, logs = load_data_with_fallback(mode)
    coin_val = sum([p.get("current_val", 0) for p in positions])
    total_portfolio = cash + coin_val

    bot_active = (bot_status == "running")

    wins = [h for h in history if (h.get("profit_usd") or 0) > 0]
    losses = [h for h in history if (h.get("profit_usd") or 0) <= 0]
    win_rate = (len(wins) / len(history) * 100) if history else 0.0
    realised_pnl = sum([h.get("profit_usd", 0) for h in history])
    unrealised_pnl = sum([p.get("pnl_usd", p.get("profit_usd", 0)) for p in positions])
    total_profit = realised_pnl + unrealised_pnl

    init_bal = 100.0 if mode == "testing" else 250.0
    net_roi = (total_profit / init_bal * 100) if init_bal > 0 else 0.0

    return jsonify({
        "mode": mode,
        "bot_status": bot_status,
        "bot_active": bot_active,
        "cash": round(cash, 2),
        "coin_value": round(coin_val, 2),
        "total_portfolio": round(total_portfolio, 2),
        "total_profit": round(total_profit, 2),
        "net_roi": round(net_roi, 2),
        "win_rate": round(win_rate, 1),
        "wins_count": len(wins),
        "losses_count": len(losses),
        "positions": positions,
        "history": history,
        "logs": logs,
        "can_buy_slots": max(0, int(cash // Config.POSITION_SIZE)),
        "can_reset": (mode == "testing"),
        "config": {
            "initial_balance": init_bal,
            "position_size": Config.POSITION_SIZE,
            "mode": mode
        }
    })


@app.route("/api/analytics")
def api_analytics():
    timeframe = request.args.get("timeframe", "all").lower()
    cash, positions, all_history, _ = load_data_with_fallback()
    coin_val = sum([p.get("current_val", 0) for p in positions])
    current_balance = cash + coin_val
    init_balance = Config.INITIAL_BALANCE

    now = time.time()
    history = all_history
    if timeframe == "today":
        history = [h for h in all_history if (h.get("closed_at") or now) >= (now - 86400)]
    elif timeframe == "week":
        history = [h for h in all_history if (h.get("closed_at") or now) >= (now - 7 * 86400)]
    elif timeframe == "month":
        history = [h for h in all_history if (h.get("closed_at") or now) >= (now - 30 * 86400)]

    total_trades = len(history)
    wins = [h for h in history if (h.get("profit_usd") or 0) > 0]
    losses = [h for h in history if (h.get("profit_usd") or 0) <= 0]
    
    win_count = len(wins)
    loss_count = len(losses)
    win_rate = (win_count / total_trades * 100) if total_trades > 0 else 0

    gross_profit = sum([h.get("profit_usd", 0) for h in wins])
    gross_loss = abs(sum([h.get("profit_usd", 0) for h in losses]))

    if timeframe == "all":
        net_profit = current_balance - init_balance
    else:
        net_profit = sum([h.get("profit_usd", 0) for h in history])
    net_roi = (net_profit / init_balance * 100) if init_balance > 0 else 0

    profit_factor = round(gross_profit / gross_loss, 2) if gross_loss > 0 else (99.9 if gross_profit > 0 else 0.0)

    avg_trade_pnl = (net_profit / total_trades) if total_trades > 0 else 0
    avg_trade_pct = (sum([h.get("pnl_pct", 0) for h in history]) / total_trades) if total_trades > 0 else 0

    best = max(history, key=lambda x: x.get("profit_usd", 0), default={})
    worst = min(history, key=lambda x: x.get("profit_usd", 0), default={})

    best_trade = {
        "symbol": best.get("symbol", "-"),
        "profit_usd": best.get("profit_usd", 0),
        "pnl_pct": best.get("pnl_pct", 0)
    }
    worst_trade = {
        "symbol": worst.get("symbol", "-"),
        "profit_usd": worst.get("profit_usd", 0),
        "pnl_pct": worst.get("pnl_pct", 0)
    }

    total_volume = sum([h.get("cost_usd", Config.POSITION_SIZE) for h in history])
    
    hold_times = []
    for h in history:
        opened = h.get("opened_at")
        closed = h.get("closed_at")
        if opened and closed and closed > opened:
            hold_times.append((closed - opened) / 60)
        else:
            hold_times.append(1.5)
    avg_hold_mins = (sum(hold_times) / len(hold_times)) if hold_times else 1.5

    tp_count = sum(1 for h in history if h.get("reason") == "TAKE_PROFIT" or (h.get("profit_usd") or 0) > 0)
    sl_count = sum(1 for h in history if h.get("reason") == "STOP_LOSS" or (h.get("profit_usd") or 0) <= 0)

    # Cumulative equity curve
    equity_curve = [{"cum_pnl": 0.0, "symbol": "START"}]
    running_pnl = 0.0
    for h in history:
        running_pnl += h.get("profit_usd", 0)
        equity_curve.append({
            "cum_pnl": round(running_pnl, 2),
            "symbol": h.get("symbol", "SOL"),
            "trade_id": h.get("id"),
            "closed_at": h.get("closed_at")
        })

    # Token stats
    token_map = {}
    for h in history:
        sym = h.get("symbol", "UNKNOWN")
        if sym not in token_map:
            token_map[sym] = {"symbol": sym, "count": 0, "wins": 0, "total_profit": 0.0}
        token_map[sym]["count"] += 1
        p_usd = h.get("profit_usd", 0)
        token_map[sym]["total_profit"] += p_usd
        if p_usd > 0:
            token_map[sym]["wins"] += 1

    token_stats = []
    for sym, t in token_map.items():
        t["win_rate"] = (t["wins"] / t["count"] * 100) if t["count"] > 0 else 0
        t["total_profit"] = round(t["total_profit"], 2)
        token_stats.append(t)
    token_stats.sort(key=lambda x: x["total_profit"], reverse=True)

    return jsonify({
        "current_balance": current_balance,
        "initial_balance": init_balance,
        "net_profit": net_profit,
        "net_roi": net_roi,
        "total_trades": total_trades,
        "win_count": win_count,
        "loss_count": loss_count,
        "win_rate": win_rate,
        "gross_profit": gross_profit,
        "gross_loss": gross_loss,
        "profit_factor": profit_factor,
        "avg_trade_pnl": avg_trade_pnl,
        "avg_trade_pct": avg_trade_pct,
        "best_trade": best_trade,
        "worst_trade": worst_trade,
        "total_volume": total_volume,
        "avg_hold_mins": avg_hold_mins,
        "tp_count": tp_count,
        "sl_count": sl_count,
        "equity_curve": equity_curve,
        "token_stats": token_stats,
        "recent_trades": list(reversed(history))[:10],
        "recent_trades_all": all_history
    })


@app.route("/api/export/transactions")
def api_export_transactions():
    filter_type = request.args.get("filter", "all")
    search_query = request.args.get("search", "").strip().lower()

    _, _, history, _ = load_data_with_fallback()

    # Filter by TP / SL
    if filter_type == "TP":
        history = [h for h in history if (h.get("profit_usd") or 0) > 0 or h.get("reason") == "TAKE_PROFIT"]
    elif filter_type == "SL":
        history = [h for h in history if (h.get("profit_usd") or 0) <= 0 or h.get("reason") == "STOP_LOSS"]

    # Filter by search query
    if search_query:
        history = [
            h for h in history
            if search_query in str(h.get("symbol", "")).lower()
            or search_query in str(h.get("address", "")).lower()
            or search_query in str(h.get("id", "")).lower()
        ]

    # Sort descending by trade ID
    sorted_history = sorted(history, key=lambda x: x.get("id", 0), reverse=True)

    rows = []
    for h in sorted_history:
        opened_ts = h.get("opened_at")
        closed_ts = h.get("closed_at")
        opened_str = time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(opened_ts)) if opened_ts else "-"
        closed_str = time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(closed_ts)) if closed_ts else "-"
        hold_sec = (closed_ts - opened_ts) if (closed_ts and opened_ts and closed_ts > opened_ts) else 0
        hold_min = round(hold_sec / 60, 1) if hold_sec > 0 else 0

        cost = h.get("cost_usd", Config.POSITION_SIZE)
        profit = h.get("profit_usd", 0)
        proceeds = h.get("proceeds", round(cost + profit, 4))
        pnl_pct = round(h.get("pnl_pct", 0), 2)
        reason = h.get("reason", "TAKE_PROFIT" if profit > 0 else "STOP_LOSS")

        rows.append({
            "Trade ID": h.get("id", "-"),
            "Symbol": f"${h.get('symbol', 'UNKNOWN')}",
            "Contract Address": h.get("address", "-"),
            "Trigger / Reason": reason,
            "Buy Price ($)": h.get("buy_price", 0),
            "Sell Price ($)": h.get("sell_price", 0),
            "Position Size ($)": cost,
            "Proceeds ($)": proceeds,
            "Net Profit/Loss ($)": profit,
            "ROI (%)": pnl_pct,
            "Hold Duration (Min)": hold_min,
            "Buy Time": opened_str,
            "Sell Time": closed_str
        })

    # Try Excel export with pandas and openpyxl, fallback to standard CSV if not installed
    try:
        import io
        import pandas as pd
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        df_trades = pd.DataFrame(rows)

        # Performance Summary Sheet Data
        total_tx = len(sorted_history)
        wins = [h for h in sorted_history if (h.get("profit_usd") or 0) > 0]
        losses = [h for h in sorted_history if (h.get("profit_usd") or 0) <= 0]
        win_cnt = len(wins)
        loss_cnt = len(losses)
        win_rate = round(win_cnt / total_tx * 100, 1) if total_tx > 0 else 0
        net_profit = sum(h.get("profit_usd", 0) for h in sorted_history)
        total_vol = sum(h.get("cost_usd", Config.POSITION_SIZE) for h in sorted_history)
        gross_win = sum(h.get("profit_usd", 0) for h in wins)
        gross_loss = abs(sum(h.get("profit_usd", 0) for h in losses))
        profit_factor = round(gross_win / gross_loss, 2) if gross_loss > 0 else (99.9 if gross_win > 0 else 0)

        summary_rows = [
            {"Metric": "Report Generated At", "Value": time.strftime('%Y-%m-%d %H:%M:%S')},
            {"Metric": "Active Filter", "Value": f"Filter: {filter_type.upper()}" + (f" | Search: '{search_query}'" if search_query else "")},
            {"Metric": "Total Trades Executed", "Value": total_tx},
            {"Metric": "Winning Trades (TP)", "Value": win_cnt},
            {"Metric": "Losing Trades (SL)", "Value": loss_cnt},
            {"Metric": "Win Rate (%)", "Value": f"{win_rate}%"},
            {"Metric": "Gross Profit ($)", "Value": f"${gross_win:.2f}"},
            {"Metric": "Gross Loss ($)", "Value": f"${gross_loss:.2f}"},
            {"Metric": "Net Realised Profit ($)", "Value": f"${net_profit:+.2f}"},
            {"Metric": "Profit Factor", "Value": profit_factor},
            {"Metric": "Total Volume Traded ($)", "Value": f"${total_vol:.2f}"},
        ]
        df_summary = pd.DataFrame(summary_rows)

        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine='openpyxl') as writer:
            df_trades.to_excel(writer, sheet_name='Trade Ledger', index=False)
            ws_trades = writer.sheets['Trade Ledger']

            df_summary.to_excel(writer, sheet_name='Performance Summary', index=False)
            ws_summary = writer.sheets['Performance Summary']

            header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            center_align = Alignment(horizontal="center", vertical="center")

            for col_idx in range(1, len(df_trades.columns) + 1):
                cell = ws_trades.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            for col in ws_trades.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws_trades.column_dimensions[col_letter].width = max(max_len + 4, 12)

            for col_idx in range(1, len(df_summary.columns) + 1):
                cell = ws_summary.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            for col in ws_summary.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws_summary.column_dimensions[col_letter].width = max(max_len + 6, 20)

        buf.seek(0)
        filename = f"Monetra_Trades_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        return send_file(
            buf,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception:
        # Fallback to CSV format using standard library
        import io
        import csv
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        mem = io.BytesIO(buf.getvalue().encode('utf-8'))
        filename = f"Monetra_Trades_{time.strftime('%Y%m%d_%H%M%S')}.csv"
        return send_file(
            mem,
            as_attachment=True,
            download_name=filename,
            mimetype="text/csv"
        )



@app.route("/api/coin/<address>")
def api_coin(address):
    _, positions, history, _ = load_data_with_fallback()

    pos = next((p for p in positions if p.get("address") == address), None)
    hist = next((h for h in reversed(history) if h.get("address") == address), None)
    
    live_data = get_token_details(address)
    
    result = {
        "address": address,
        "symbol": "Unknown",
        "current_price": 0,
    }
    
    if live_data:
        result.update({
            "pair_address": live_data.get("pair_address"),
            "current_price": live_data.get("price_usd", 0),
            "symbol": live_data.get("symbol", "UNKNOWN"),
            "dex_url": live_data.get("dex_url"),
            "volume_24h": live_data.get("volume_24h"),
            "liquidity_usd": live_data.get("liquidity_usd"),
        })
        
    if pos:
        result.update(pos)
        if live_data:
            result["current_val"] = pos.get("tokens_count", 0) * live_data.get("price_usd", 0)
        result["status"] = "ACTIVE"
    elif hist:
        result.update(hist)
        result["status"] = "CLOSED"
        if not live_data:
            result["pair_address"] = address
    else:
        result["status"] = "NOT_OWNED"
        if not live_data:
            result["pair_address"] = address
        
    return jsonify(result)


@app.route("/api/control/<action>", methods=["POST"])
def api_control(action):
    state = get_bot_state()
    current_mode = state.get("mode", "testing")

    if action == "start":
        state["bot_status"] = "running"
        save_bot_state(state)
        try:
            subprocess.run(["systemctl", "start", "trading-bot"])
        except Exception:
            pass
        log_event("STATUS", f"Bot diaktifkan (Running) pada Mode {current_mode.upper()}. Memindai token Solana...")
        return jsonify({"success": True, "bot_status": "running", "message": "Bot sniper berhasil dimulai!"})

    elif action == "pause":
        state["bot_status"] = "paused"
        save_bot_state(state)
        log_event("STATUS", f"Bot dijeda (Paused) pada Mode {current_mode.upper()}. Pembelian baru ditahan, posisi aktif tetap dipantau.")
        return jsonify({"success": True, "bot_status": "paused", "message": "Bot berhasil dijeda!"})

    elif action == "stop":
        state["bot_status"] = "stopped"
        save_bot_state(state)
        try:
            subprocess.run(["systemctl", "stop", "trading-bot"])
        except Exception:
            pass
        log_event("STATUS", f"Bot dihentikan (Stopped) pada Mode {current_mode.upper()}. Semua pemindaian dihentikan.")
        return jsonify({"success": True, "bot_status": "stopped", "message": "Bot berhasil dihentikan!"})

    elif action == "mode":
        req_data = request.get_json(silent=True) or {}
        new_mode = req_data.get("mode") or request.args.get("mode", "testing")
        if new_mode in ["testing", "production"]:
            state["mode"] = new_mode
            save_bot_state(state)
            log_event("STATUS", f"Beralih ke MODE: {new_mode.upper()}. Saldo dan riwayat dialihkan.")
            return jsonify({"success": True, "mode": new_mode, "message": f"Berhasil beralih ke Mode {new_mode.title()}!"})
        return jsonify({"error": "Mode tidak valid"}), 400

    elif action == "reset":
        if current_mode == "production":
            return jsonify({
                "success": False,
                "error": "Akses Ditolak: Reset portofolio dinonaktifkan di Mode Production demi melindungi saldo dan aset nyata Solana Anda!"
            }), 403

        # Testing Mode Reset
        default_test_cash = 100.00
        try:
            reset_db(default_test_cash)
        except Exception:
            pass
        try:
            with open(getattr(Config, "DATA_FILE", "portfolio.json"), "w") as f:
                json.dump({"cash": default_test_cash, "positions": []}, f, indent=2)
            with open(getattr(Config, "HISTORY_FILE", "trade_history.json"), "w") as f:
                json.dump([], f, indent=2)
        except Exception:
            pass
        log_event("STATUS", f"Portofolio Mode Testing berhasil direset ke modal awal ${default_test_cash:.2f}.")
        return jsonify({
            "success": True,
            "message": f"Portofolio Testing berhasil direset ke ${default_test_cash:.2f}!"
        })

    else:
        return jsonify({"error": "Unknown action"}), 400


@app.route("/api/control/sell/<address>", methods=["POST"])
def api_control_sell(address):
    try:
        positions = get_positions()
        found = False
        for p in positions:
            if p["address"] == address:
                p["target_sl_price"] = 9999999999
                p["target_tp_price"] = 0
                found = True
        if found:
            save_positions(positions)
            with open(Config.DATA_FILE, "w") as f:
                json.dump({"cash": get_wallet(), "positions": positions}, f, indent=2)
            return jsonify({"success": True, "message": "Perintah jual dikirim! Bot akan menutup posisi dalam 15 detik."})
        return jsonify({"error": "Posisi tidak ditemukan"}), 404
    except Exception as e:
        return jsonify({"error": f"Gagal update posisi: {e}"}), 500


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5050)
